from __future__ import annotations

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


DOCX_PATH = Path(
    r"C:\Users\Admin\Desktop\DoANTotNghiep\pod-ai-evaluator\docs\plantuml\BAO_CAO_DO_AN_Ky.docx"
)
BACKUP_PATH = DOCX_PATH.with_name("BAO_CAO_DO_AN_Ky.backup_before_chapter3.docx")
OUTPUT_PATH = DOCX_PATH.with_name("BAO_CAO_DO_AN_Ky_Chuong3.docx")


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        p.remove(child)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def set_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "D9D9D9")


def add_paragraph(doc: Document, text: str = "", style: str | None = None, bold_label: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_label and text.startswith(bold_label):
        run = paragraph.add_run(bold_label)
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        rest = text[len(bold_label) :]
        if rest:
            run = paragraph.add_run(rest)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
    else:
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
    return paragraph


def add_heading(doc: Document, text: str, level: int):
    style_name = f"Heading {level}" if level > 0 else "Title"
    if style_name not in [style.name for style in doc.styles]:
        style_name = "Heading 2" if "Heading 2" in [style.name for style in doc.styles] else "Normal"
    paragraph = doc.add_paragraph(style=style_name)
    run = paragraph.add_run(text)
    run.bold = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.color.rgb = None
    return paragraph


def add_image_note(doc: Document, image_name: str, caption: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"[Chèn hình {image_name} tại đây]")
    run.bold = True
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    run = cap.add_run(caption)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def add_bullets(doc: Document, items: list[str]):
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(f"- {item}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_numbered(doc: Document, items: list[str]):
    for index, item in enumerate(items, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(f"{index}. {item}")
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_spec_block(
    doc: Document,
    title: str,
    actor: str,
    description: str,
    condition: str,
    main_flow: list[str],
    alternative_flow: list[str] | None = None,
    result: str | None = None,
):
    add_heading(doc, title, 3)
    for label, value in [
        ("Tác nhân: ", actor),
        ("Mô tả: ", description),
        ("Điều kiện: ", condition),
    ]:
        add_paragraph(doc, label + value, bold_label=label)

    add_paragraph(doc, "Luồng sự kiện chính:", bold_label="Luồng sự kiện chính:")
    add_numbered(doc, main_flow)

    if alternative_flow:
        add_paragraph(doc, "Luồng sự kiện rẽ nhánh:", bold_label="Luồng sự kiện rẽ nhánh:")
        add_bullets(doc, alternative_flow)

    if result:
        add_paragraph(doc, "Kết quả: " + result, bold_label="Kết quả: ")


def add_collection_table(doc: Document):
    table = doc.add_table(rows=1, cols=3)
    set_table_borders(table)
    headers = ["Collection", "Vai trò", "Dữ liệu chính"]
    for idx, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], "1F4E5F")
        set_cell_text(table.rows[0].cells[idx], header, bold=True)

    rows = [
        ("platforms", "Lưu thông tin sàn", "code, name"),
        ("products", "Lưu thông tin sản phẩm", "platform_code, external_product_id, name, image_urls"),
        ("reviews", "Lưu review thô và review đã chuẩn hóa", "rating, content, review_time, review_images"),
        ("review_aspects", "Lưu khía cạnh được phân tích từ nội dung", "aspect_name, sentiment, confidence"),
        ("review_trust_analyses", "Lưu điểm tin cậy và gợi ý xử lý", "trust_score, trust_label, seller_action"),
        ("ingestion_jobs", "Theo dõi tiến trình kéo dữ liệu", "status, total_reviews, created_at"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)


def remove_existing_chapter3(doc: Document):
    start = None
    for idx, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip().lower()
        if text.startswith("chương 3") or text.startswith("chuong 3"):
            start = idx
            break
    if start is None:
        return
    for paragraph in list(doc.paragraphs[start + 1 :]):
        delete_paragraph(paragraph)
    clear_paragraph(doc.paragraphs[start])
    doc.paragraphs[start].style = doc.styles["Heading 1"]
    run = doc.paragraphs[start].add_run("CHƯƠNG 3: PHÂN TÍCH THIẾT KẾ HỆ THỐNG")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)


def build_chapter3(doc: Document):
    add_heading(doc, "3.1 Tổng quan chương", 2)
    add_paragraph(
        doc,
        "Chương này trình bày quá trình phân tích và thiết kế hệ thống trí tuệ nhân tạo hỗ trợ seller phân tích phản hồi khách hàng trên các sàn thương mại điện tử. Nội dung chương tập trung vào việc xác định tác nhân, chức năng, luồng xử lý nghiệp vụ, kiến trúc thành phần và mô hình dữ liệu của hệ thống.",
    )
    add_paragraph(
        doc,
        "Các sơ đồ được sử dụng trong chương gồm use case, phân rã chức năng, activity, sequence, component, state, deployment, class diagram và ERD. Ở mỗi vị trí có ghi chú chèn hình, người viết có thể xuất ảnh từ PlantUML rồi dán vào đúng vị trí tương ứng.",
    )

    add_heading(doc, "3.2 Biểu đồ use case tổng quát", 2)
    add_image_note(
        doc,
        "FINAL_01_use_case_tong_quat",
        "Hình 3.1. Biểu đồ use case tổng quát của hệ thống",
    )
    add_paragraph(
        doc,
        "Biểu đồ use case tổng quát cho thấy hệ thống có ba nhóm tác nhân chính: Seller, Admin và Sàn thương mại điện tử. Seller là người sử dụng trực tiếp hệ thống, thực hiện các thao tác dán link sản phẩm, mở phiên sàn khi cần xác minh, kéo dữ liệu review, xem dashboard, lọc review theo số sao và xem các review cần chú ý.",
    )
    add_paragraph(
        doc,
        "Admin có vai trò theo dõi log và cấu hình hệ thống. Các chức năng này hỗ trợ quá trình vận hành, đặc biệt trong trường hợp API của sàn thay đổi, phát sinh lỗi khi thu thập dữ liệu hoặc cần điều chỉnh số trang, thời gian chờ và phiên trình duyệt. Sàn thương mại điện tử đóng vai trò nguồn dữ liệu, cung cấp thông tin sản phẩm, review và hình ảnh phục vụ phân tích.",
    )

    add_heading(doc, "3.3 Biểu đồ phân rã chức năng", 2)
    add_image_note(
        doc,
        "FINAL_02_phan_ra_chuc_nang",
        "Hình 3.2. Biểu đồ phân rã chức năng của hệ thống",
    )
    add_paragraph(
        doc,
        "Hệ thống được chia thành năm nhóm chức năng chính. Nhóm xử lý link sản phẩm chịu trách nhiệm nhận link, nhận diện sàn và tách mã sản phẩm. Nhóm thu thập review thực hiện việc gọi API hoặc mở phiên trình duyệt, sau đó lấy review theo số sao và từng trang. Nhóm phân tích review đảm nhận phân tích nội dung, kiểm định hình ảnh và tính điểm độ tin cậy.",
    )
    add_paragraph(
        doc,
        "Nhóm lưu trữ dữ liệu đảm bảo toàn bộ dữ liệu sản phẩm, review, kết quả phân tích và trạng thái tiến trình được lưu vào MongoDB. Nhóm dashboard cho seller hiển thị kết quả phân tích, cho phép lọc review theo số sao và ưu tiên các review cần xử lý.",
    )

    add_heading(doc, "3.4 Đặc tả các chức năng chính", 2)
    add_image_note(
        doc,
        "FINAL_03_use_case_nhap_link",
        "Hình 3.3. Use case nhập và xử lý link sản phẩm",
    )
    add_spec_block(
        doc,
        "3.4.1 Đặc tả use case nhập và xử lý link sản phẩm",
        "Seller.",
        "Cho phép seller dán đường dẫn sản phẩm từ Tiki, Lazada hoặc TikTok Shop để hệ thống nhận diện sàn, tách mã sản phẩm và kiểm tra dữ liệu đã có.",
        "Seller phải nhập một đường dẫn sản phẩm thuộc sàn được hệ thống hỗ trợ.",
        [
            "Seller copy đường dẫn sản phẩm từ trình duyệt.",
            "Seller dán đường dẫn vào ô nhập link trên giao diện.",
            "Seller nhấn nút Đọc link.",
            "Frontend gửi link đến backend.",
            "Backend kiểm tra định dạng URL và nhận diện sàn thương mại điện tử.",
            "Backend tách mã sản phẩm từ đường dẫn.",
            "Hệ thống lấy thông tin sản phẩm và ảnh gốc nếu có.",
            "Hệ thống kiểm tra trong MongoDB xem sản phẩm đã có dữ liệu review hay chưa.",
            "Frontend hiển thị sàn, mã sản phẩm, số review đã có và ảnh gốc sản phẩm.",
        ],
        [
            "A1: Nếu link không hợp lệ, hệ thống hiển thị thông báo yêu cầu seller kiểm tra lại đường dẫn.",
            "A2: Nếu chưa lấy được ảnh gốc sản phẩm, hệ thống vẫn cho phép kéo review nhưng ghi nhận trạng thái thiếu ảnh gốc.",
        ],
        "Seller biết được hệ thống đã nhận diện đúng sản phẩm và có thể tiếp tục kéo dữ liệu review.",
    )

    add_image_note(
        doc,
        "FINAL_04_use_case_thu_thap_va_phan_tich",
        "Hình 3.4. Use case thu thập và phân tích review",
    )
    add_spec_block(
        doc,
        "3.4.2 Đặc tả use case thu thập và phân tích review",
        "Seller, System, Sàn thương mại điện tử.",
        "Cho phép hệ thống thu thập review của sản phẩm, chuẩn hóa dữ liệu, phân tích nội dung, kiểm định hình ảnh và tính điểm độ tin cậy.",
        "Hệ thống đã nhận diện được sàn và mã sản phẩm. Nếu sàn yêu cầu xác minh, seller cần mở phiên trình duyệt hợp lệ.",
        [
            "Seller nhấn nút Kéo dữ liệu.",
            "Backend tạo tiến trình thu thập dữ liệu.",
            "Hệ thống xác định phương thức lấy dữ liệu phù hợp với từng sàn.",
            "Nếu cần, hệ thống sử dụng phiên trình duyệt Playwright để lấy cookie hoặc token hợp lệ.",
            "Crawler lấy review theo từng mức sao và từng trang.",
            "Hệ thống chuẩn hóa dữ liệu review gồm mã review, số sao, nội dung, thời gian và ảnh review.",
            "Dữ liệu review được lưu vào MongoDB.",
            "Module phân tích nội dung đánh giá từ khóa, mức tiêu cực và dấu hiệu cần chú ý.",
            "Module phân tích ảnh kiểm tra mức phù hợp giữa ảnh review và ảnh gốc sản phẩm.",
            "Hệ thống tổng hợp kết quả và tính điểm độ tin cậy.",
        ],
        [
            "A1: Nếu API bị giới hạn hoặc yêu cầu captcha, hệ thống yêu cầu seller mở phiên sàn.",
            "A2: Nếu một mức sao không còn review, hệ thống chuyển sang mức sao tiếp theo.",
            "A3: Nếu review không có ảnh, hệ thống bỏ qua bước so sánh ảnh và vẫn phân tích bằng nội dung.",
        ],
        "Review được lưu và phân tích, sẵn sàng hiển thị trên dashboard.",
    )

    add_image_note(
        doc,
        "FINAL_05_use_case_dashboard",
        "Hình 3.5. Use case dashboard phân tích cho seller",
    )
    add_spec_block(
        doc,
        "3.4.3 Đặc tả use case dashboard phân tích",
        "Seller.",
        "Cho phép seller xem kết quả phân tích review, lọc review theo số sao, xem bằng chứng hình ảnh và gợi ý hành động.",
        "Sản phẩm đã có dữ liệu review trong MongoDB hoặc đã được hệ thống kéo dữ liệu thành công.",
        [
            "Seller mở dashboard của sản phẩm.",
            "Frontend gọi API lấy thống kê tổng quan.",
            "Frontend gọi API lấy danh sách review cần chú ý.",
            "Frontend gọi API lấy kết quả so sánh hình ảnh.",
            "Hệ thống hiển thị tổng số review, điểm tin cậy trung bình và trạng thái bằng chứng ảnh.",
            "Seller chọn bộ lọc theo số sao nếu cần.",
            "Hệ thống cập nhật bảng review theo bộ lọc.",
            "Seller xem từng review, điểm độ tin cậy, mức bằng chứng và gợi ý hành động.",
        ],
        [
            "A1: Nếu chưa có review trong MongoDB, dashboard hiển thị thông báo cần kéo dữ liệu trước.",
            "A2: Nếu chưa có kết quả phân tích ảnh, hệ thống hiển thị trạng thái chưa chạy phân tích ảnh.",
        ],
        "Seller xác định được các review cần ưu tiên đọc và xử lý.",
    )

    add_heading(doc, "3.5 Biểu đồ hoạt động", 2)
    add_image_note(
        doc,
        "FINAL_06_activity_tong_quat",
        "Hình 3.6. Biểu đồ hoạt động tổng quát từ link sản phẩm đến dashboard",
    )
    add_paragraph(
        doc,
        "Biểu đồ hoạt động tổng quát mô tả luồng xử lý chính của hệ thống. Quy trình bắt đầu khi seller dán link sản phẩm và nhấn Đọc link. Frontend gửi dữ liệu đến backend, backend kiểm tra URL, tách mã sản phẩm và kiểm tra dữ liệu trong MongoDB. Nếu dữ liệu đã tồn tại, hệ thống có thể hiển thị dashboard ngay. Nếu chưa có dữ liệu, hệ thống thực hiện thu thập review từ sàn.",
    )
    add_paragraph(
        doc,
        "Sau khi có review, hệ thống thực hiện phân tích nội dung, kiểm định hình ảnh, tính điểm tin cậy và gắn nhãn review cần chú ý. Cuối cùng, frontend tải dữ liệu summary, review và bằng chứng ảnh để hiển thị dashboard cho seller.",
    )

    add_image_note(
        doc,
        "FINAL_19_activity_thu_thap_chi_tiet",
        "Hình 3.7. Biểu đồ hoạt động chi tiết chức năng thu thập review",
    )
    add_paragraph(
        doc,
        "Biểu đồ này làm rõ quá trình thu thập review. Hệ thống nhận platform_code và product_id, tạo ingestion job và xác định phương thức thu thập. Nếu sàn yêu cầu phiên trình duyệt, hệ thống kiểm tra session trước khi gửi request. Trong quá trình lấy dữ liệu, hệ thống duyệt từng mức sao và từng trang, sau đó chuẩn hóa dữ liệu, loại bỏ trùng lặp và lưu vào MongoDB.",
    )

    add_image_note(
        doc,
        "FINAL_20_activity_phan_tich_chi_tiet",
        "Hình 3.8. Biểu đồ hoạt động chi tiết chức năng phân tích review",
    )
    add_paragraph(
        doc,
        "Biểu đồ hoạt động phân tích review thể hiện cách hệ thống xử lý từng đánh giá sau khi đã được lưu. Nếu review có nội dung, hệ thống phân tích từ khóa, cảm xúc và nhóm vấn đề. Nếu review có ảnh, hệ thống lấy ảnh review và ảnh gốc sản phẩm để so sánh. Sau đó, các điểm thành phần được tổng hợp thành trust_score và chuyển thành nhãn độ tin cậy cùng gợi ý hành động cho seller.",
    )

    add_heading(doc, "3.6 Biểu đồ tuần tự", 2)
    sequence_notes = [
        (
            "FINAL_07_sequence_keo_du_lieu",
            "Hình 3.9. Biểu đồ tuần tự tổng quát quá trình kéo dữ liệu",
            "Biểu đồ tuần tự tổng quát cho thấy sự phối hợp giữa Seller, Frontend, Backend, Crawler, AI Analysis, MongoDB và Sàn thương mại điện tử. Seller thao tác trên frontend, backend điều phối crawler để lấy review, lưu dữ liệu vào MongoDB và gọi module AI để phân tích trước khi trả kết quả về dashboard.",
        ),
        (
            "FINAL_11_sequence_doc_link_san_pham",
            "Hình 3.10. Biểu đồ tuần tự chức năng đọc link sản phẩm",
            "Ở chức năng đọc link, frontend gửi URL đến backend. Backend gọi product resolver để nhận diện sàn và mã sản phẩm, sau đó gọi metadata service để lấy tên sản phẩm và ảnh gốc. Kết quả được đối chiếu với MongoDB để xác định sản phẩm đã có review hay chưa.",
        ),
        (
            "FINAL_12_sequence_mo_phien_san",
            "Hình 3.11. Biểu đồ tuần tự chức năng mở phiên sàn",
            "Biểu đồ này mô tả trường hợp sàn yêu cầu đăng nhập hoặc captcha. Backend khởi tạo browser session bằng Playwright, mở trang sản phẩm và hiển thị trình duyệt để seller xác minh. Khi phiên hợp lệ, hệ thống có thể dùng cookie hoặc token của phiên đó để thu thập dữ liệu.",
        ),
        (
            "FINAL_13_sequence_thu_thap_review",
            "Hình 3.12. Biểu đồ tuần tự thu thập review theo số sao và trang",
            "Trong quá trình thu thập, ingestion service tạo job và duyệt review theo từng mức sao, từng trang. Mỗi lần crawler lấy được dữ liệu, review sẽ được chuẩn hóa và upsert vào MongoDB. Khi hoàn tất, job được cập nhật trạng thái completed.",
        ),
        (
            "FINAL_14_sequence_phan_tich_text",
            "Hình 3.13. Biểu đồ tuần tự phân tích nội dung review",
            "Biểu đồ này mô tả quá trình phân tích text. Service lấy danh sách review từ MongoDB, chuyển từng review sang module phát hiện rủi ro nội dung để trích xuất từ khóa, mức tiêu cực và loại phàn nàn. Kết quả được lưu vào review_aspects.",
        ),
        (
            "FINAL_15_sequence_kiem_dinh_hinh_anh",
            "Hình 3.14. Biểu đồ tuần tự kiểm định bằng chứng hình ảnh",
            "Với các review có ảnh, hệ thống tải ảnh review và ảnh gốc sản phẩm từ nguồn ảnh. Image Evidence Module so sánh ảnh review với ảnh sản phẩm, đồng thời đối chiếu ảnh với tên sản phẩm và nội dung review. Kết quả cuối cùng là điểm phù hợp hình ảnh.",
        ),
        (
            "FINAL_16_sequence_tinh_diem_tin_cay",
            "Hình 3.15. Biểu đồ tuần tự tính điểm độ tin cậy review",
            "Điểm độ tin cậy được tính bằng cách tổng hợp nhiều yếu tố gồm số sao, nội dung bình luận và bằng chứng ảnh. Review Trust Module trả về trust_score, trust_label và seller_action. Kết quả được lưu trong collection review_trust_analyses.",
        ),
        (
            "FINAL_17_sequence_xem_dashboard",
            "Hình 3.16. Biểu đồ tuần tự xem dashboard phân tích",
            "Khi seller mở dashboard, frontend gọi ba nhóm API chính: summary, reviews cần chú ý và image comparisons. Backend đọc dữ liệu từ MongoDB rồi trả kết quả về frontend để hiển thị các khối thống kê và bảng review.",
        ),
        (
            "FINAL_18_sequence_loc_review_theo_sao",
            "Hình 3.17. Biểu đồ tuần tự lọc review theo số sao",
            "Khi seller chọn bộ lọc sao, frontend gửi tham số rating đến backend. Backend query MongoDB theo platform, product_id và rating, sau đó sắp xếp kết quả theo thời gian mới nhất rồi trả lại danh sách review phù hợp.",
        ),
    ]
    for image_name, caption, explanation in sequence_notes:
        add_image_note(doc, image_name, caption)
        add_paragraph(doc, explanation)

    add_heading(doc, "3.7 Thiết kế kiến trúc và triển khai", 2)
    add_image_note(
        doc,
        "FINAL_08_component_kien_truc",
        "Hình 3.18. Biểu đồ component kiến trúc hệ thống",
    )
    add_paragraph(
        doc,
        "Kiến trúc hệ thống gồm Frontend, Backend API, Crawler, AI Analysis và MongoDB. Frontend được xây dựng bằng React và TypeScript, cung cấp giao diện cho seller. Backend API được xây dựng bằng FastAPI, đóng vai trò điều phối nghiệp vụ. Crawler giao tiếp với các sàn thương mại điện tử để lấy review. AI Analysis xử lý nội dung, hình ảnh và tính điểm độ tin cậy. MongoDB lưu toàn bộ dữ liệu sản phẩm, review, kết quả phân tích và trạng thái job.",
    )

    add_image_note(
        doc,
        "FINAL_21_state_ingestion_job",
        "Hình 3.19. Biểu đồ trạng thái ingestion job",
    )
    add_paragraph(
        doc,
        "Biểu đồ trạng thái mô tả vòng đời của một tiến trình thu thập dữ liệu. Job bắt đầu ở trạng thái Created, sau đó chuyển sang ResolvingProduct để xử lý link. Nếu cần xác minh, job chuyển sang WaitingSession. Khi có thể thu thập, job chuyển sang Crawling, Saving và Analyzing. Nếu hoàn tất, job ở trạng thái Completed. Nếu phát sinh lỗi link, API, MongoDB hoặc module AI, job chuyển sang Failed.",
    )

    add_image_note(
        doc,
        "FINAL_22_deployment_trien_khai",
        "Hình 3.20. Biểu đồ triển khai hệ thống",
    )
    add_paragraph(
        doc,
        "Biểu đồ triển khai thể hiện cách các thành phần được vận hành trong môi trường thực tế. Người dùng truy cập frontend bằng trình duyệt web. Frontend gọi REST API đến backend FastAPI chạy bằng Uvicorn. Backend kết nối MongoDB, gọi các module AI và sử dụng Playwright Browser khi cần mở phiên trình duyệt để lấy dữ liệu từ sàn thương mại điện tử.",
    )

    add_heading(doc, "3.8 Thiết kế lớp và cơ sở dữ liệu", 2)
    add_image_note(
        doc,
        "FINAL_09_class_domain_model",
        "Hình 3.21. Biểu đồ lớp mô hình miền dữ liệu",
    )
    add_paragraph(
        doc,
        "Biểu đồ lớp mô tả các thực thể chính trong miền dữ liệu của hệ thống. Platform đại diện cho sàn thương mại điện tử. Product lưu thông tin sản phẩm, bao gồm mã sản phẩm, tên, URL và ảnh gốc. Review lưu thông tin đánh giá của khách hàng, gồm số sao, nội dung, thời gian và ảnh review. ReviewAspect lưu các khía cạnh được phân tích từ nội dung. ReviewTrustAnalysis lưu điểm tin cậy, nhãn độ tin cậy, mức bằng chứng và gợi ý hành động cho seller. IngestionJob ghi nhận quá trình kéo dữ liệu.",
    )

    add_image_note(
        doc,
        "FINAL_10_erd_mongodb",
        "Hình 3.22. Sơ đồ ERD MongoDB",
    )
    add_paragraph(
        doc,
        "Cơ sở dữ liệu của hệ thống sử dụng MongoDB với các collection chính như platforms, products, reviews, review_aspects, review_trust_analyses và ingestion_jobs. Các collection này được thiết kế để phục vụ cả hai mục tiêu: lưu dữ liệu review đã thu thập và lưu kết quả phân tích phục vụ dashboard.",
    )
    add_collection_table(doc)
    add_paragraph(
        doc,
        "Thiết kế dữ liệu theo hướng tách collection giúp hệ thống dễ mở rộng khi bổ sung thêm sàn thương mại điện tử mới hoặc thêm module phân tích mới. Ví dụ, nếu cần bổ sung phân tích xu hướng theo thời gian, hệ thống có thể tạo thêm collection hoặc trường dữ liệu mới mà không làm thay đổi cấu trúc cốt lõi của sản phẩm và review.",
    )

    add_heading(doc, "3.9 Kết luận chương", 2)
    add_paragraph(
        doc,
        "Chương 3 đã trình bày quá trình phân tích và thiết kế hệ thống hỗ trợ seller phân tích phản hồi khách hàng. Thông qua các biểu đồ use case, phân rã chức năng, activity, sequence, component, state, deployment, class diagram và ERD, có thể thấy hệ thống được thiết kế theo hướng tách biệt rõ ràng giữa giao diện người dùng, xử lý nghiệp vụ, thu thập dữ liệu, phân tích AI và lưu trữ dữ liệu.",
    )
    add_paragraph(
        doc,
        "Thiết kế này phù hợp với mục tiêu của đề tài là giúp seller giảm thời gian đọc review thủ công, ưu tiên các đánh giá có rủi ro, kiểm tra bằng chứng hình ảnh và đưa ra quyết định xử lý dựa trên dữ liệu. Bên cạnh đó, kiến trúc hệ thống vẫn có khả năng mở rộng thêm sàn thương mại điện tử mới, bổ sung module phân tích mới hoặc phát triển chức năng tự động gợi ý phản hồi trong các giai đoạn tiếp theo.",
    )


def normalize_styles(doc: Document):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.name is None:
                run.font.name = "Times New Roman"
            if run.font.size is None and paragraph.style.name.startswith("Heading"):
                continue
            if run.font.size is None:
                run.font.size = Pt(12)


def main():
    if not BACKUP_PATH.exists():
        BACKUP_PATH.write_bytes(DOCX_PATH.read_bytes())

    doc = Document(DOCX_PATH)
    remove_existing_chapter3(doc)
    build_chapter3(doc)
    normalize_styles(doc)
    try:
        doc.save(DOCX_PATH)
    except PermissionError:
        doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
