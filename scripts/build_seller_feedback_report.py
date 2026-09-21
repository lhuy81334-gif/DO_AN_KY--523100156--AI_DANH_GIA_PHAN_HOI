import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT_DIR = Path(__file__).resolve().parents[1]
SOURCE_DOCX = ROOT_DIR.parent / "docs" / "BÁO CÁO PHÂN TÍCH ĐỀ TÀI.docx"
OUTPUT_DOCX = ROOT_DIR / "docs" / "Bao_cao_phan_tich_de_tai_seller_feedback.docx"
MEDIA_DIR = ROOT_DIR / "tmp_doc_review" / "report_media"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, bold=False, size=12, color="000000") -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, alignment=None, space_after=6, line_spacing=1.15) -> None:
    if alignment is not None:
        paragraph.alignment = alignment
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    for run in paragraph.runs:
        if run.text:
            set_font(run)


def add_paragraph(doc, text="", style=None, alignment=None, bold=False, size=12, space_after=6):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_font(run, bold=bold, size=size)
    set_paragraph(paragraph, alignment=alignment, space_after=space_after)
    return paragraph


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run)
        set_paragraph(paragraph, space_after=3)


def add_heading(doc, text: str, level: int) -> None:
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        set_font(run, bold=True, size=15 if level == 1 else 13)
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        set_font(run, bold=True, size=11, color="FFFFFF")
        set_cell_shading(cell, "1F4E79")
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            cell.width = Cm(widths[i])

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            set_font(run, size=10.5)
            set_cell_borders(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if len(value) > 25 else WD_ALIGN_PARAGRAPH.CENTER
            if widths:
                cells[i].width = Cm(widths[i])

    add_paragraph(doc, "", space_after=4)
    return table


def add_caption(doc, text: str) -> None:
    paragraph = add_paragraph(doc, text, alignment=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=8)
    for run in paragraph.runs:
        run.italic = True


def add_use_case_spec(doc, caption: str, rows: list[list[str]]) -> None:
    add_table(doc, ["Thành phần", "Nội dung đặc tả"], rows, [4.2, 11.8])
    add_caption(doc, caption)


def extract_media() -> list[Path]:
    MEDIA_DIR.mkdir(parents=True, exist_ok=True)
    if not SOURCE_DOCX.exists():
        return []
    media_paths = []
    with zipfile.ZipFile(SOURCE_DOCX) as archive:
        for name in archive.namelist():
            if name.startswith("word/media/"):
                target = MEDIA_DIR / Path(name).name
                target.write_bytes(archive.read(name))
                media_paths.append(target)
    return sorted(media_paths)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True


def add_cover(doc: Document) -> None:
    for text in [
        "TRƯỜNG ĐẠI HỌC PHƯƠNG ĐÔNG",
        "KHOA CÔNG NGHỆ SỐ VÀ TRUYỀN THÔNG",
        "NGÀNH CÔNG NGHỆ THÔNG TIN",
    ]:
        add_paragraph(doc, text, alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=3)
    add_paragraph(doc, "_____________", alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=28)
    add_paragraph(doc, "BÁO CÁO PHÂN TÍCH ĐỀ TÀI", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=18, space_after=18)
    add_paragraph(
        doc,
        "ĐỀ TÀI: XÂY DỰNG HỆ THỐNG TRÍ TUỆ NHÂN TẠO HỖ TRỢ SELLER PHÂN TÍCH VÀ KIỂM ĐỊNH ĐÁNH GIÁ SẢN PHẨM",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        bold=True,
        size=14,
        space_after=28,
    )
    for text in [
        "Sinh viên thực hiện: Lương Trọng Huy",
        "MSSV: 523100156",
        "Ngày hoàn thành: 12/09/2026",
        "Công nghệ chính: ReactJS, FastAPI, MongoDB, CLIP/OpenCLIP",
    ]:
        add_paragraph(doc, text, alignment=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=6)
    add_paragraph(doc, "Hà Nội, 2026", alignment=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12, space_after=0)
    doc.add_page_break()


def add_static_toc(doc: Document) -> None:
    add_heading(doc, "MỤC LỤC", 1)
    toc_lines = [
        "LỜI MỞ ĐẦU",
        "CHƯƠNG 1. TỔNG QUAN VỀ ĐỀ TÀI",
        "1.1. Bài toán thực tế và lý do chọn đề tài",
        "1.2. Mục tiêu của đề tài",
        "1.3. Đối tượng sử dụng, phạm vi và giới hạn",
        "1.4. Dữ liệu đầu vào, đầu ra và các ràng buộc chính",
        "1.5. Công nghệ và công cụ sử dụng",
        "CHƯƠNG 2. PHÂN TÍCH VÀ THIẾT KẾ GIẢI PHÁP",
        "2.1. Phân tích yêu cầu chức năng và phi chức năng",
        "2.2. Đặc tả use case",
        "2.3. Kiến trúc hệ thống",
        "2.4. Thiết kế dữ liệu",
        "2.5. Thiết kế giao diện và luồng xử lý",
        "2.6. Phương pháp đánh giá độ tin cậy review",
        "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN",
    ]
    for line in toc_lines:
        add_paragraph(doc, line, space_after=2)
    doc.add_page_break()


def add_intro(doc: Document) -> None:
    add_heading(doc, "LỜI MỞ ĐẦU", 1)
    add_paragraph(
        doc,
        "Trong hoạt động kinh doanh trên các sàn thương mại điện tử, đánh giá của khách hàng là một nguồn dữ liệu quan trọng ảnh hưởng trực tiếp đến uy tín sản phẩm và quyết định mua hàng của người tiêu dùng. Tuy nhiên, seller thường phải đọc thủ công một lượng lớn review trên nhiều sàn khác nhau, trong đó có cả review chỉ có sao, review nội dung mơ hồ, review có ảnh chưa chắc phù hợp với sản phẩm, hoặc review tiêu cực cần được xử lý sớm.",
    )
    add_paragraph(
        doc,
        "Đề tài này xây dựng một hệ thống hỗ trợ seller thu thập, chuẩn hóa và phân tích đánh giá sản phẩm từ Tiki và Lazada. Hệ thống cho phép người dùng dán link sản phẩm, tự động nhận diện sàn, kéo dữ liệu review qua API, lưu vào MongoDB, phân tích nội dung chữ, bằng chứng hình ảnh và độ tin cậy tổng thể. Kết quả được trình bày trên dashboard để seller ưu tiên đọc những đánh giá có dấu hiệu bất thường hoặc cần kiểm tra trước.",
    )
    add_paragraph(
        doc,
        "Do phạm vi đồ án tập trung vào phân tích và hỗ trợ ra quyết định, hệ thống không thay seller kết luận tuyệt đối review thật hay giả. Thay vào đó, hệ thống đóng vai trò như một lớp sàng lọc thông minh, giúp seller tiết kiệm thời gian và có thêm căn cứ khi xử lý phản hồi khách hàng.",
    )
    doc.add_page_break()


def add_chapter_1(doc: Document) -> None:
    add_heading(doc, "CHƯƠNG 1. TỔNG QUAN VỀ ĐỀ TÀI", 1)
    add_heading(doc, "1.1. Bài toán thực tế và lý do chọn đề tài", 2)
    add_paragraph(
        doc,
        "Seller trên các sàn như Tiki và Lazada thường nhận được nhiều đánh giá khác nhau về sản phẩm, chất lượng đóng gói, vận chuyển và trải nghiệm sử dụng. Việc đọc toàn bộ review bằng tay tốn nhiều thời gian, đặc biệt khi seller cần phân biệt đâu là phản hồi có giá trị, đâu là review có dấu hiệu bất thường cần kiểm tra thêm.",
    )
    add_paragraph(
        doc,
        "Một số trường hợp gây khó khăn gồm: đánh giá sao thấp nhưng nội dung không rõ lý do, đánh giá sao cao nhưng lại phàn nàn trong nội dung, ảnh review không khớp với sản phẩm đã mua, hoặc review cũ không còn phản ánh tình trạng hiện tại của shop. Vì vậy, cần có một hệ thống hỗ trợ thu thập và phân tích review theo hướng tự động, minh bạch và dễ sử dụng cho seller.",
    )
    add_table(
        doc,
        ["Tiêu chí", "Xử lý thủ công", "Seller Center", "Hệ thống đề tài"],
        [
            ["Nguồn dữ liệu", "Từng sàn riêng lẻ", "Trong từng sàn", "Tiki và Lazada trong một dashboard"],
            ["Cách tìm review cần chú ý", "Đọc thủ công", "Lọc theo sao/ngày", "Tự xếp hạng review có rủi ro"],
            ["Kiểm tra hình ảnh", "Nhìn bằng mắt", "Hiển thị ảnh review", "So ảnh review với sản phẩm và nội dung"],
            ["Gợi ý xử lý", "Tự soạn", "Phụ thuộc seller", "Sinh phản hồi gợi ý để seller chỉnh/copy"],
        ],
        [3.2, 4.0, 4.0, 5.0],
    )
    add_caption(doc, "Bảng 1.1. So sánh cách xử lý đánh giá sản phẩm")

    add_heading(doc, "1.2. Mục tiêu của đề tài", 2)
    add_heading(doc, "1.2.1. Mục tiêu tổng quát", 3)
    add_paragraph(
        doc,
        "Xây dựng hệ thống trí tuệ nhân tạo hỗ trợ seller thu thập, phân tích và kiểm định độ tin cậy của đánh giá sản phẩm trên các sàn thương mại điện tử, từ đó giúp seller ưu tiên những review cần kiểm tra và khai thác insight khách hàng hiệu quả hơn.",
    )
    add_heading(doc, "1.2.2. Mục tiêu cụ thể", 3)
    add_bullets(
        doc,
        [
            "Xây dựng giao diện cho phép seller dán link sản phẩm Tiki hoặc Lazada để bắt đầu phân tích.",
            "Tự động nhận diện sàn, bóc mã sản phẩm và gọi API review tương ứng.",
            "Chuẩn hóa dữ liệu review, hình ảnh, thời gian comment và lưu vào MongoDB.",
            "Phân tích nội dung review, số sao và mức độ cụ thể của phản hồi.",
            "Kiểm định ảnh review với thông tin sản phẩm và nội dung người mua viết.",
            "Tính điểm độ tin cậy, phát hiện review bất thường và hiển thị dashboard cho seller.",
        ],
    )

    add_heading(doc, "1.3. Đối tượng sử dụng, phạm vi và giới hạn", 2)
    add_paragraph(
        doc,
        "Đối tượng sử dụng chính là seller hoặc người vận hành shop trên sàn thương mại điện tử. Người dùng không cần thao tác với mã nguồn; họ chỉ cần dán link sản phẩm, chờ hệ thống thu thập dữ liệu và xem danh sách review cần chú ý.",
    )
    add_paragraph(
        doc,
        "Phạm vi hiện tại tập trung vào hai nguồn dữ liệu là Tiki và Lazada. Hệ thống ưu tiên phân tích review có nội dung chữ và review có hình ảnh. Các chức năng tự động phản hồi trực tiếp lên Seller Center chưa được triển khai trong phạm vi hiện tại do cần tài khoản seller, quyền API chính thức và cơ chế xác nhận của người bán.",
    )
    add_bullets(
        doc,
        [
            "Hệ thống không khẳng định tuyệt đối review thật hay giả, mà đưa ra điểm tin cậy và dấu hiệu cần kiểm tra.",
            "Dữ liệu Lazada có thể phụ thuộc cookie/token hợp lệ do cơ chế chống bot của nền tảng.",
            "Một số review chỉ có sao hoặc không có nội dung sẽ có giá trị phân tích thấp.",
            "Kết quả so sánh ảnh sử dụng mô hình AI/ML nên cần được seller kiểm chứng lại trong các trường hợp quan trọng.",
        ],
    )

    add_heading(doc, "1.4. Dữ liệu đầu vào, đầu ra và các ràng buộc chính", 2)
    add_table(
        doc,
        ["Nhóm dữ liệu", "Mô tả"],
        [
            ["Đầu vào", "Link sản phẩm Tiki/Lazada, dữ liệu review từ API, ảnh review, ảnh sản phẩm, số sao và thời gian comment."],
            ["Đầu ra", "Dashboard phân tích, danh sách review cần kiểm tra, điểm độ tin cậy, trạng thái ảnh, gợi ý phản hồi cho seller."],
            ["Ràng buộc", "API có thể thay đổi, token Lazada có thể hết hạn, dữ liệu review có thể không đầy đủ so với tổng số đánh giá hiển thị trên sàn."],
        ],
        [4.0, 11.5],
    )
    add_caption(doc, "Bảng 1.2. Dữ liệu đầu vào, đầu ra và ràng buộc")

    add_heading(doc, "1.5. Công nghệ và công cụ sử dụng", 2)
    add_table(
        doc,
        ["Nhóm công nghệ", "Công nghệ sử dụng", "Vai trò"],
        [
            ["Front-end", "ReactJS, TypeScript, Vite, CSS/TailwindCSS", "Xây dựng dashboard và giao diện dán link sản phẩm."],
            ["Back-end", "Python, FastAPI, Uvicorn, Pydantic", "Xây dựng API, xử lý request và điều phối pipeline phân tích."],
            ["Database", "MongoDB, MongoDB Compass", "Lưu sản phẩm, review, ảnh, kết quả phân tích và lịch sử import."],
            ["AI/ML", "review_trust.py, image_evidence.py, CLIP/OpenCLIP", "Phân tích độ tin cậy và kiểm định bằng chứng hình ảnh."],
            ["Thu thập dữ liệu", "REST API, requests, cURL, DevTools Network Inspector", "Phân tích request của sàn và lấy dữ liệu review."],
            ["Kiểm thử", "Pytest, TypeScript compiler", "Kiểm tra API, service và tính đúng kiểu dữ liệu front-end."],
        ],
        [3.0, 5.0, 7.5],
    )
    add_caption(doc, "Bảng 1.3. Công nghệ và công cụ sử dụng")


def add_chapter_2(doc: Document, media_paths: list[Path]) -> None:
    doc.add_page_break()
    add_heading(doc, "CHƯƠNG 2. PHÂN TÍCH VÀ THIẾT KẾ GIẢI PHÁP", 1)
    add_heading(doc, "2.1. Phân tích yêu cầu chức năng và phi chức năng", 2)
    add_table(
        doc,
        ["Mã", "Yêu cầu chức năng", "Mô tả"],
        [
            ["FR01", "Nhập link sản phẩm", "Seller dán link Tiki/Lazada để hệ thống nhận diện sàn và mã sản phẩm."],
            ["FR02", "Thu thập đánh giá", "Hệ thống gọi API review, lấy nội dung, số sao, thời gian và ảnh."],
            ["FR03", "Chuẩn hóa dữ liệu", "Chuyển response từng sàn về schema chung và lưu MongoDB."],
            ["FR04", "Phân tích độ tin cậy", "Tính trust score dựa trên sao, nội dung, ảnh, thời gian và dấu hiệu bất thường."],
            ["FR05", "Kiểm định ảnh", "So ảnh review với ảnh sản phẩm và nội dung người mua viết."],
            ["FR06", "Dashboard phân tích", "Hiển thị tổng quan và danh sách review cần seller đọc trước."],
            ["FR07", "Xem chi tiết review", "Hiển thị nội dung, ngày comment, ảnh, độ tin cậy và link sản phẩm."],
            ["FR08", "Gợi ý phản hồi", "Sinh phản hồi mẫu để seller chỉnh sửa và copy khi cần."],
        ],
        [2.0, 4.0, 9.5],
    )
    add_caption(doc, "Bảng 2.1. Yêu cầu chức năng")
    add_table(
        doc,
        ["Mã", "Yêu cầu phi chức năng", "Mô tả"],
        [
            ["NFR01", "Dễ sử dụng", "Seller thao tác bằng giao diện dán link, không cần chạy lệnh thủ công."],
            ["NFR02", "Minh bạch", "Mỗi kết quả phân tích phải có điểm số, trạng thái và dấu hiệu đi kèm."],
            ["NFR03", "Mở rộng", "Có thể bổ sung thêm sàn khác như Shopee bằng connector mới."],
            ["NFR04", "An toàn dữ liệu", "Không lưu mật khẩu seller; token/cookie nếu dùng chỉ phục vụ thu thập dữ liệu trong môi trường kiểm thử."],
            ["NFR05", "Hiệu năng chấp nhận được", "Dữ liệu sau khi đã có trong Mongo phải tải dashboard nhanh; phân tích ảnh có thể chạy riêng do tốn tài nguyên."],
        ],
        [2.0, 4.0, 9.5],
    )
    add_caption(doc, "Bảng 2.2. Yêu cầu phi chức năng")

    add_heading(doc, "2.2. Đặc tả use case", 2)
    add_paragraph(
        doc,
        "Phần use case được tách theo bốn sơ đồ chính. Sơ đồ tổng quát mô tả phạm vi toàn hệ thống, còn ba sơ đồ chi tiết được phân rã thành từng bảng đặc tả riêng cho từng use case chính. Cách trình bày này giúp người đọc thấy rõ tác nhân nào thực hiện chức năng nào, điều kiện bắt đầu ra sao và hệ thống trả kết quả như thế nào.",
    )

    if media_paths:
        doc.add_picture(str(media_paths[0]), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_caption(doc, "Hình 2.1. Sơ đồ use case tổng quát của hệ thống")
    add_table(
        doc,
        ["Nhóm", "Tác nhân", "Use case chính", "Ý nghĩa"],
        [
            ["Nhóm Seller", "Seller", "Nhập link dẫn, xem dashboard phân tích, tra cứu review nghi ngờ", "Cho phép seller bắt đầu phân tích sản phẩm và ưu tiên đọc các review cần chú ý."],
            ["Nhóm Server", "Server", "Thu thập đánh giá, kiểm định bằng chứng hình ảnh, phát hiện bất thường", "Tự động xử lý dữ liệu phía sau, từ gọi API đến tính điểm tin cậy."],
            ["Nhóm Admin", "Admin", "Quản lí tài khoản và phân quyền, quản lí API", "Quản trị người dùng và cấu hình phục vụ việc thu thập dữ liệu."],
        ],
        [3.0, 2.5, 5.0, 6.0],
    )
    add_caption(doc, "Bảng 2.3. Đặc tả use case tổng quát")

    if len(media_paths) > 1:
        doc.add_picture(str(media_paths[1]), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_caption(doc, "Hình 2.2. Sơ đồ use case chi tiết phía Seller")

    add_use_case_spec(
        doc,
        "Bảng 2.4. Đặc tả use case Nhập link dẫn",
        [
            ["Mã use case", "UC-S01"],
            ["Tên use case", "Nhập link dẫn"],
            ["Tác nhân chính", "Seller"],
            ["Mục tiêu", "Seller cung cấp link sản phẩm cần phân tích để hệ thống tự nhận diện sàn và tạo tiến trình thu thập review."],
            ["Tiền điều kiện", "Seller đang ở dashboard; backend, MongoDB và cấu hình kết nối API đang hoạt động."],
            ["Quan hệ include", "Kiểm tra hợp lệ đường dẫn; tạo tiến trình."],
            ["Luồng chính", "Seller dán link sản phẩm. Hệ thống kiểm tra định dạng link, xác định sàn Tiki hoặc Lazada, bóc product_id/item_id, tạo ingestion job và chuyển sang bước thu thập đánh giá."],
            ["Luồng thay thế", "Nếu link đã từng được thu thập, hệ thống có thể dùng dữ liệu hiện có và cho phép seller tải lại dữ liệu mới khi cần."],
            ["Ngoại lệ", "Link không thuộc sàn hỗ trợ, thiếu mã sản phẩm, API thay đổi hoặc token Lazada hết hạn."],
            ["Hậu điều kiện", "Một tiến trình thu thập review được tạo, trạng thái xử lý được lưu vào ingestion_jobs và dashboard có thể cập nhật theo kết quả mới."],
        ],
    )
    add_use_case_spec(
        doc,
        "Bảng 2.5. Đặc tả use case Xem Dashboard phân tích",
        [
            ["Mã use case", "UC-S02"],
            ["Tên use case", "Xem Dashboard phân tích"],
            ["Tác nhân chính", "Seller"],
            ["Mục tiêu", "Seller xem nhanh tình hình đánh giá của sản phẩm để biết tổng review, review có ảnh, review cần kiểm tra và độ tin cậy trung bình."],
            ["Tiền điều kiện", "Sản phẩm đã có dữ liệu review trong MongoDB hoặc vừa hoàn tất tiến trình thu thập."],
            ["Quan hệ include", "Hiển thị thống kê tổng quan."],
            ["Quan hệ extend", "Lọc và phân loại theo sàn; cảnh báo chỉ số rủi ro."],
            ["Luồng chính", "Seller chọn sàn và sản phẩm. Front-end gọi API thống kê, API danh sách review và API so sánh ảnh. Hệ thống hiển thị các chỉ số tổng quan, bảng review cần chú ý và kết quả kiểm định ảnh."],
            ["Luồng thay thế", "Seller có thể lọc theo Tiki/Lazada, theo sản phẩm, theo review có ảnh hoặc theo review bị gắn cờ."],
            ["Ngoại lệ", "Chưa có dữ liệu cho sản phẩm, MongoDB chưa kết nối hoặc pipeline phân tích ảnh chưa chạy xong."],
            ["Hậu điều kiện", "Seller nắm được nhóm review cần ưu tiên đọc và có thể mở chi tiết từng review."],
        ],
    )
    add_use_case_spec(
        doc,
        "Bảng 2.6. Đặc tả use case Tra cứu review nghi ngờ",
        [
            ["Mã use case", "UC-S03"],
            ["Tên use case", "Tra cứu review nghi ngờ"],
            ["Tác nhân chính", "Seller"],
            ["Mục tiêu", "Seller tìm và xem các review có dấu hiệu bất thường như sao thấp, nội dung mâu thuẫn, ảnh không khớp sản phẩm hoặc bằng chứng yếu."],
            ["Tiền điều kiện", "Review đã được chuẩn hóa và có kết quả trust_analysis hoặc image_evidence."],
            ["Quan hệ extend", "Tìm kiếm và lọc; xem chi tiết hồ sơ đánh giá."],
            ["Luồng chính", "Seller mở danh sách review nghi ngờ. Hệ thống sắp xếp theo mức độ cần kiểm tra và thời gian comment gần nhất. Seller bấm vào mã review để xem chi tiết nội dung, ảnh, điểm tin cậy và lý do bị gắn cờ."],
            ["Luồng thay thế", "Seller lọc riêng review từ 1 đến 3 sao, review có ảnh, review ảnh lệch sản phẩm hoặc review có rating-text mismatch."],
            ["Ngoại lệ", "Review không có nội dung chữ, không có ảnh hoặc ảnh không tải được từ nguồn sàn."],
            ["Hậu điều kiện", "Seller xác định được review nào cần đọc kỹ, phản hồi, theo dõi hoặc kiểm tra thêm trong Seller Center."],
        ],
    )

    if len(media_paths) > 2:
        doc.add_picture(str(media_paths[2]), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_caption(doc, "Hình 2.3. Sơ đồ use case chi tiết phía Admin")
    add_use_case_spec(
        doc,
        "Bảng 2.7. Đặc tả use case Quản lí tài khoản và phân quyền",
        [
            ["Mã use case", "UC-A01"],
            ["Tên use case", "Quản lí tài khoản và phân quyền"],
            ["Tác nhân chính", "Admin"],
            ["Mục tiêu", "Admin kiểm soát danh sách người dùng và quyền truy cập vào các chức năng quản trị, dashboard và cấu hình dữ liệu."],
            ["Tiền điều kiện", "Admin đã đăng nhập bằng tài khoản có quyền quản trị."],
            ["Quan hệ include", "Quản lý danh sách người dùng; phân quyền."],
            ["Luồng chính", "Admin mở màn hình quản trị, xem danh sách tài khoản, thêm hoặc cập nhật người dùng, gán vai trò phù hợp và lưu thay đổi."],
            ["Luồng thay thế", "Admin có thể khóa tài khoản không còn sử dụng hoặc điều chỉnh quyền khi người dùng đổi vai trò."],
            ["Ngoại lệ", "Thông tin tài khoản không hợp lệ, trùng email/tên đăng nhập hoặc admin không đủ quyền."],
            ["Hậu điều kiện", "Quyền truy cập của người dùng được cập nhật và hệ thống ghi nhận trạng thái mới."],
        ],
    )
    add_use_case_spec(
        doc,
        "Bảng 2.8. Đặc tả use case Quản lí API",
        [
            ["Mã use case", "UC-A02"],
            ["Tên use case", "Quản lí API"],
            ["Tác nhân chính", "Admin"],
            ["Mục tiêu", "Admin theo dõi và cập nhật cấu hình kết nối đến nguồn dữ liệu Tiki/Lazada để hệ thống thu thập review ổn định."],
            ["Tiền điều kiện", "Admin có quyền quản trị cấu hình hệ thống."],
            ["Quan hệ include", "Theo dõi trạng thái."],
            ["Luồng chính", "Admin xem trạng thái API, kiểm tra token/cookie, cập nhật cấu hình khi nguồn dữ liệu thay đổi và kiểm tra kết quả kết nối."],
            ["Luồng thay thế", "Nếu token Lazada hết hạn, admin cập nhật lại cookie hoặc chuyển hệ thống sang trạng thái yêu cầu cấu hình mới."],
            ["Ngoại lệ", "API bị đổi cấu trúc, request bị chặn, cookie thiếu quyền hoặc cấu hình sai định dạng."],
            ["Hậu điều kiện", "Trạng thái nguồn dữ liệu được cập nhật để server biết connector nào sẵn sàng hoạt động."],
        ],
    )

    if len(media_paths) > 3:
        doc.add_picture(str(media_paths[3]), width=Inches(5.8))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_caption(doc, "Hình 2.4. Sơ đồ use case chi tiết phía Server")
    add_use_case_spec(
        doc,
        "Bảng 2.9. Đặc tả use case Thu thập đánh giá",
        [
            ["Mã use case", "UC-SV01"],
            ["Tên use case", "Thu thập đánh giá"],
            ["Tác nhân chính", "Server"],
            ["Mục tiêu", "Tự động gọi API của sàn để lấy review sản phẩm và lưu dữ liệu vào MongoDB theo schema chung."],
            ["Tiền điều kiện", "Có mã sản phẩm hợp lệ, connector tương ứng đã được cấu hình và nguồn API có thể truy cập."],
            ["Quan hệ include", "Kết nối API và xử lý; lưu trữ và khởi tạo view."],
            ["Luồng chính", "Server nhận ingestion job, xác định connector Tiki/Lazada, gọi API theo từng trang hoặc từng nhóm sao, nhận response, chuẩn hóa review, lưu products/reviews và cập nhật các view phục vụ dashboard."],
            ["Luồng thay thế", "Nếu sản phẩm đã có dữ liệu, server cập nhật thêm review mới hoặc bỏ qua review trùng dựa trên external_review_id."],
            ["Ngoại lệ", "API trả lỗi, hết token, request timeout, response rỗng hoặc dữ liệu thiếu trường bắt buộc."],
            ["Hậu điều kiện", "Review được lưu trong MongoDB, ingestion job có trạng thái hoàn tất hoặc lỗi kèm lý do."],
        ],
    )
    add_use_case_spec(
        doc,
        "Bảng 2.10. Đặc tả use case Kiểm định bằng chứng hình ảnh",
        [
            ["Mã use case", "UC-SV02"],
            ["Tên use case", "Kiểm định bằng chứng hình ảnh"],
            ["Tác nhân chính", "Server"],
            ["Mục tiêu", "Đánh giá ảnh review có phù hợp với sản phẩm đã mua và có hỗ trợ nội dung khách hàng viết hay không."],
            ["Tiền điều kiện", "Review có ảnh hoặc có đường dẫn ảnh; sản phẩm có ảnh chuẩn/mô tả sản phẩm để đối chiếu."],
            ["Quan hệ include", "Tải và lưu trữ ảnh; so khớp ảnh."],
            ["Quan hệ extend", "Phát hiện hình ảnh không hợp lệ."],
            ["Luồng chính", "Server tải ảnh review, lấy ảnh/mô tả sản phẩm, trích đặc trưng bằng mô hình thị giác, so điểm ảnh với ảnh sản phẩm, so ảnh với nội dung review và sinh image_evidence_score."],
            ["Luồng thay thế", "Nếu ảnh không tải được, hệ thống đánh dấu chưa phân tích ảnh và giảm độ mạnh của bằng chứng thay vì loại bỏ toàn bộ review."],
            ["Ngoại lệ", "Ảnh lỗi, link ảnh hết hạn, ảnh quá mờ, ảnh không liên quan hoặc mô hình không đủ tự tin."],
            ["Hậu điều kiện", "Review có trạng thái ảnh như phù hợp, có thể phù hợp, không rõ hoặc có dấu hiệu không khớp sản phẩm."],
        ],
    )
    add_use_case_spec(
        doc,
        "Bảng 2.11. Đặc tả use case Phát hiện bất thường",
        [
            ["Mã use case", "UC-SV03"],
            ["Tên use case", "Phát hiện bất thường"],
            ["Tác nhân chính", "Server"],
            ["Mục tiêu", "Tự động tìm review cần seller kiểm tra dựa trên số sao, nội dung, ảnh, thời gian và điểm tin cậy."],
            ["Tiền điều kiện", "Review đã được lưu và có dữ liệu phân tích cơ bản về nội dung hoặc hình ảnh."],
            ["Quan hệ include", "Tính toán điểm tin cậy."],
            ["Luồng chính", "Server tính rating_sentiment, text_sentiment, sentiment_gap, content_specificity_score, image_evidence_score và review_trust_score. Sau đó hệ thống gắn flags và xếp loại review theo mức độ cần seller chú ý."],
            ["Luồng thay thế", "Review 4 đến 5 sao có ảnh không khớp chỉ được ghi nhận để tham khảo, còn review từ 1 đến 3 sao được ưu tiên kiểm tra kỹ hơn."],
            ["Ngoại lệ", "Review thiếu nội dung chữ, thiếu ảnh hoặc dữ liệu thời gian không rõ ràng khiến điểm tin cậy chỉ mang tính tham khảo."],
            ["Hậu điều kiện", "Dashboard có danh sách review nghi ngờ, review có ảnh cần kiểm tra và gợi ý phản hồi để seller xử lý."],
        ],
    )

    add_heading(doc, "2.3. Kiến trúc hệ thống", 2)
    add_paragraph(
        doc,
        "Hệ thống được tổ chức theo kiến trúc ba lớp: giao diện người dùng, dịch vụ back-end và lớp dữ liệu/phân tích. Front-end ReactJS cung cấp dashboard cho seller. Back-end FastAPI tiếp nhận link, điều phối connector Tiki/Lazada, chuẩn hóa dữ liệu và trả kết quả phân tích. MongoDB lưu dữ liệu sản phẩm, review, ảnh, kết quả trust analysis và lịch sử import.",
    )
    add_table(
        doc,
        ["Lớp", "Thành phần", "Vai trò"],
        [
            ["Giao diện", "ReactJS, TypeScript, Vite", "Dán link, xem dashboard, mở chi tiết review và copy phản hồi gợi ý."],
            ["Dịch vụ", "FastAPI, Pydantic, Uvicorn", "Cung cấp REST API, validate dữ liệu, điều phối pipeline."],
            ["Dữ liệu", "MongoDB", "Lưu products, reviews, review_aspects, review_trust_analyses và ingestion_jobs."],
            ["AI/ML", "review_trust.py, image_evidence.py", "Phân tích nội dung, số sao, hình ảnh và độ tin cậy."],
            ["Nguồn ngoài", "Tiki API, Lazada MTOP API", "Cung cấp dữ liệu đánh giá sản phẩm."],
        ],
        [3.0, 5.0, 7.5],
    )
    add_caption(doc, "Bảng 2.12. Kiến trúc thành phần hệ thống")

    add_heading(doc, "2.4. Thiết kế dữ liệu", 2)
    add_table(
        doc,
        ["Collection", "Vai trò dữ liệu chính"],
        [
            ["platforms", "Lưu thông tin sàn thương mại điện tử như Tiki, Lazada."],
            ["products", "Lưu sản phẩm, mã sản phẩm ngoài sàn, link sản phẩm, ảnh sản phẩm và thông tin mô tả."],
            ["reviews", "Lưu review đã chuẩn hóa gồm số sao, nội dung, ngày comment, ảnh review, source và trust_analysis."],
            ["review_aspects", "Lưu các khía cạnh được bóc tách từ nội dung review như giao hàng, đóng gói, chất lượng."],
            ["review_trust_analyses", "Lưu kết quả phân tích độ tin cậy, flags và điểm đánh giá."],
            ["ingestion_jobs", "Theo dõi lịch sử import hoặc kéo dữ liệu review."],
            ["counters", "Sinh mã tăng dần nội bộ cho các bản ghi."],
        ],
        [4.0, 11.5],
    )
    add_caption(doc, "Bảng 2.13. Thiết kế collection trong MongoDB")

    add_heading(doc, "2.5. Thiết kế giao diện và luồng xử lý", 2)
    add_paragraph(
        doc,
        "Giao diện chính là Review Evidence Dashboard. Seller có thể dán link sản phẩm hoặc nhập trực tiếp sàn và ID sản phẩm. Sau khi dữ liệu được tải, dashboard hiển thị các chỉ số tổng quan như tổng review, review có ảnh, review cần kiểm tra, trust score trung bình, đồng thời cung cấp bảng review cần seller chú ý và bảng so sánh ảnh.",
    )
    add_bullets(
        doc,
        [
            "Bước 1: Seller dán link sản phẩm.",
            "Bước 2: Backend nhận diện sàn và mã sản phẩm.",
            "Bước 3: Backend gọi API review, chuẩn hóa và lưu MongoDB.",
            "Bước 4: Module AI/ML phân tích độ tin cậy và bằng chứng ảnh.",
            "Bước 5: Dashboard hiển thị review cần kiểm tra theo mức độ ưu tiên và thời gian comment.",
        ],
    )

    add_heading(doc, "2.6. Phương pháp đánh giá độ tin cậy review", 2)
    add_paragraph(
        doc,
        "Hệ thống đánh giá review theo nhiều tín hiệu thay vì chỉ dựa vào số sao. Các tín hiệu chính gồm: mức sao, nội dung chữ, độ cụ thể của phản hồi, sự khớp giữa số sao và nội dung, bằng chứng ảnh, độ khớp giữa ảnh review với ảnh sản phẩm, và thời điểm comment.",
    )
    add_table(
        doc,
        ["Tín hiệu", "Cách sử dụng trong hệ thống"],
        [
            ["Số sao", "Review từ 1 đến 3 sao được ưu tiên kiểm tra kỹ hơn; review 4 đến 5 sao chủ yếu dùng lấy insight."],
            ["Nội dung chữ", "Review quá ngắn hoặc không rõ lý do bị giảm trọng số bằng chứng."],
            ["Sao và nội dung", "Nếu sao cao nhưng nội dung phàn nàn, hoặc sao thấp nhưng nội dung tích cực, hệ thống gắn cờ mismatch."],
            ["Ảnh review", "Ảnh được so với mô tả sản phẩm, ảnh sản phẩm và claim trong nội dung review."],
            ["Thời gian comment", "Review gần hiện tại được ưu tiên hơn vì phản ánh tình trạng sản phẩm/vận hành mới hơn."],
            ["Trust score", "Điểm tổng hợp từ các tín hiệu để phân loại mức tin cậy: tốt, trung bình, yếu hoặc rất yếu."],
        ],
        [4.0, 11.5],
    )
    add_caption(doc, "Bảng 2.14. Các tín hiệu dùng để đánh giá độ tin cậy review")


def add_conclusion(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", 1)
    add_paragraph(
        doc,
        "Đề tài đã xây dựng được hướng tiếp cận phù hợp cho bài toán hỗ trợ seller xử lý đánh giá khách hàng trên sàn thương mại điện tử. Thay vì chỉ hiển thị danh sách review, hệ thống bổ sung lớp phân tích giúp phát hiện review cần chú ý, đánh giá mức độ tin cậy và kiểm tra bằng chứng hình ảnh.",
    )
    add_paragraph(
        doc,
        "Kết quả hiện tại cho phép seller dán link sản phẩm, thu thập dữ liệu từ Tiki/Lazada, lưu dữ liệu vào MongoDB, phân tích nội dung và hình ảnh, sau đó xem dashboard để ưu tiên review cần kiểm tra. Hệ thống vẫn cần tiếp tục hoàn thiện về khả năng tự động phản hồi qua Seller Center, mở rộng thêm nguồn dữ liệu Shopee và cải thiện độ chính xác của mô hình so sánh ảnh trong các trường hợp phức tạp.",
    )
    add_bullets(
        doc,
        [
            "Tích hợp API chính thức của Seller Center nếu có quyền truy cập hợp lệ.",
            "Bổ sung cơ chế human-in-the-loop để seller duyệt phản hồi trước khi gửi.",
            "Huấn luyện hoặc tinh chỉnh mô hình thị giác trên bộ ảnh review thực tế lớn hơn.",
            "Thêm báo cáo xu hướng theo thời gian và phân tích nguyên nhân phàn nàn theo từng nhóm sản phẩm.",
        ],
    )


def main() -> None:
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    media_paths = extract_media()
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_static_toc(doc)
    add_intro(doc)
    add_chapter_1(doc)
    add_chapter_2(doc, media_paths)
    add_conclusion(doc)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
